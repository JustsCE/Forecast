"""Workbook — NORDE-forecast.xlsx.

Summary : the Horizon P&L matrix at company level, Jan t .. Dec t+1 with FY columns.
Per line: months down, column blocks across — TOTAL | TOTAL division | entity … — each block Amount | LY | YoY %.
Blue cells = forecast months. Totals are SUM formulas. YoY is a formula with red/green conditional formatting.
"""
import re
from collections import defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col

from actuals import DIVISIONS, ENTITIES
from forecast import current_month, forecast, months, shift

BLUE = PatternFill("solid", fgColor="DCE6F1")
GREY = PatternFill("solid", fgColor="F2F2F2")
RED, RED_FONT = PatternFill("solid", fgColor="FFC7CE"), Font(color="9C0006")
GREEN, GREEN_FONT = PatternFill("solid", fgColor="C6EFCE"), Font(color="006100")
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP, BOTTOM = Border(top=THIN), Border(bottom=THIN)
BOLD = Font(bold=True)
INT, PCT = "#,##0", "0.0%"
TOTALS = [("1.99999 KOPĀ IEŅĒMUMI", "1."), ("2.99999 KOPĀ IZDEVUMI", "2."), ("3.99999 PEĻŅA", None)]


def put(ws, r, c, v, fmt=None, bold=False, fill=None, border=None):
    cell = ws.cell(row=r, column=c, value=v)
    if fmt:
        cell.number_format = fmt
    if bold:
        cell.font = BOLD
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    return cell


def plus(cols, r):
    return "=" + "+".join(f"{col(c)}{r}" for c in cols)


def yoy(c, r):
    return f"=IF({col(c + 1)}{r}=0,0,{col(c)}{r}/{col(c + 1)}{r}-1)"


def yoy_rule(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["-0.1"], fill=RED, font=RED_FONT))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0.1"], fill=GREEN, font=GREEN_FONT))


def build(actuals, subs):
    current = current_month(actuals)
    fc = forecast(actuals, subs)
    value = lambda e, l, m: actuals.get((e, l, m), 0) if m < current else fc.get((e, l, m), 0)
    carried = defaultdict(set)
    for e, l, _ in list(actuals) + list(subs):
        carried[l].add(e)
    lines = sorted(carried)

    wb = Workbook()
    summary(wb.active, lines, carried, value, current)
    for line in lines:
        sheet(wb.create_sheet(re.sub(r"[:\\/?*\[\]]", " ", line)[:31]), line, sorted(carried[line]), value, current)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def summary(ws, lines, carried, value, current):
    """A line | B FY t-1 | C..N Jan..Dec t | O FY t | P FY % | Q..AB Jan..Dec t+1 | AC FY t+1 | AD FY %"""
    ws.title = "Summary"
    t = int(current[:4])
    put(ws, 1, 1, f"NORDE forecast · actuals through {current} (MTD) · forecast from {current}", bold=True)
    heads = [f"FY {t - 1}"] + months(t) + [f"FY {t}", "FY %"] + months(t + 1) + [f"FY {t + 1}", "FY %"]
    put(ws, 2, 1, "P&L line", bold=True, border=BOTTOM)
    for c, h in enumerate(heads, 2):
        put(ws, 2, c, h, bold=True, border=BOTTOM)
    month_cols = list(zip(months(t), range(3, 15))) + list(zip(months(t + 1), range(17, 29)))
    first = 3
    for r, line in enumerate(lines, first):
        put(ws, r, 1, line)
        put(ws, r, 2, sum(value(e, line, m) for e in carried[line] for m in months(t - 1)), INT)
        for m, c in month_cols:
            fc = m >= current
            put(ws, r, c, sum(value(e, line, m) for e in carried[line]), INT, fill=BLUE if fc else None, border=BOX if fc else None)
    last = first + len(lines) - 1
    for k, (name, prefix) in enumerate(TOTALS):
        r = last + 1 + k
        rows = [i for i, l in enumerate(lines, first) if l.startswith(prefix)] if prefix else [last + 1, last + 2]
        put(ws, r, 1, name, bold=True, border=TOP)
        for c in [2] + [c for _, c in month_cols]:
            put(ws, r, c, "=" + "+".join(f"{col(c)}{i}" for i in rows), INT, bold=True, border=TOP)
    for r in range(first, last + 4):
        put(ws, r, 15, f"=SUM(C{r}:N{r})", INT, bold=True)
        put(ws, r, 16, f"=IF(B{r}=0,0,O{r}/B{r}-1)", PCT)
        put(ws, r, 29, f"=SUM(Q{r}:AB{r})", INT, bold=True)
        put(ws, r, 30, f"=IF(O{r}=0,0,AC{r}/O{r}-1)", PCT)
    yoy_rule(ws, f"P{first}:P{last + 3}")
    yoy_rule(ws, f"AD{first}:AD{last + 3}")
    ws.column_dimensions["A"].width = 48
    for c in range(2, 31):
        ws.column_dimensions[col(c)].width = 12
    ws.freeze_panes = "B3"
    ws.sheet_view.zoomScale = 70


def layout(entities):
    """Column blocks [(label, col, kind)] and {division col: [entity cols]}; a grey gap column precedes every division."""
    blocks, groups, c = [("TOTAL", 2, "company")], {}, 6
    for d in DIVISIONS:
        ents = [e for e in entities if ENTITIES[e][1] == d]
        if not ents:
            continue
        blocks.append((f"TOTAL {d}", c, "division"))
        groups[c] = [c + 3 * (i + 1) for i in range(len(ents))]
        for e, ec in zip(ents, groups[c]):
            blocks.append((f"{e} {ENTITIES[e][0]}", ec, e))
        c = groups[c][-1] + 4
    return blocks, groups


def sheet(ws, line, entities, value, current):
    """A month | B:D TOTAL | gap | TOTAL division | entity … — rows 3..26 = Jan t .. Dec t+1, then TOTAL t, TOTAL t+1."""
    t = int(current[:4])
    span = months(t) + months(t + 1)
    first, last = 3, 2 + len(span)
    blocks, groups = layout(entities)

    put(ws, 1, 1, line, bold=True)
    put(ws, 2, 1, "Month", bold=True, border=BOTTOM)
    for r, m in enumerate(span, first):
        put(ws, r, 1, m, bold=True)
    put(ws, last + 1, 1, f"TOTAL {t}", bold=True, border=TOP)
    put(ws, last + 2, 1, f"TOTAL {t + 1}", bold=True, border=TOP)
    for g in [c - 1 for c in groups]:
        for r in range(1, last + 3):
            put(ws, r, g, None, fill=GREY)
        ws.column_dimensions[col(g)].width = 3

    for label, c, kind in blocks:
        ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + 2)
        put(ws, 1, c, label, bold=True).alignment = Alignment(horizontal="center")
        for i, h in enumerate(("Amount", "LY", "YoY %")):
            put(ws, 2, c + i, h, bold=True, border=BOTTOM)
            ws.column_dimensions[col(c + i)].width = 12
        for r, m in enumerate(span, first):
            if kind == "company":
                a, ly = plus(list(groups), r), plus([d + 1 for d in groups], r)
            elif kind == "division":
                a, ly = plus(groups[c], r), plus([e + 1 for e in groups[c]], r)
            else:
                a, ly = value(kind, line, m), value(kind, line, shift(m, -12))
            fc = m >= current
            put(ws, r, c, a, INT, fill=BLUE if fc else None, border=BOX if fc else None)
            put(ws, r, c + 1, ly, INT)
            put(ws, r, c + 2, yoy(c, r), PCT)
        for r, (lo, hi) in ((last + 1, (first, first + 11)), (last + 2, (first + 12, last))):
            put(ws, r, c, f"=SUM({col(c)}{lo}:{col(c)}{hi})", INT, bold=True, border=TOP)
            put(ws, r, c + 1, f"=SUM({col(c + 1)}{lo}:{col(c + 1)}{hi})", INT, bold=True, border=TOP)
            put(ws, r, c + 2, yoy(c, r), PCT, bold=True, border=TOP)
        yoy_rule(ws, f"{col(c + 2)}{first}:{col(c + 2)}{last}")
    for ents in groups.values():
        for e in ents:
            for i in range(3):
                ws.column_dimensions[col(e + i)].outlineLevel = 1
                ws.column_dimensions[col(e + i)].hidden = True
    ws.column_dimensions["A"].width = 12
    ws.freeze_panes = "B3"
    ws.sheet_view.zoomScale = 70
