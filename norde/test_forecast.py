"""Synthetic ledger: 3 entities, 2 lines, 2025-01..2026-10; 2026 runs +10 % on 2025, so every YoY term is 0.1."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

import actuals as A
import forecast as F
import workbook as W

REV, EXP = "1.1 Auto pārdošanas ieņēmumi", "2.1 Auto iegādes izdevumi"
ENTS = ["02.1", "02.2", "03.2"]


@pytest.fixture
def actuals():
    out = {}
    for y, base in ((2025, 100.0), (2026, 110.0)):
        for m in F.months(y):
            if m > "2026-10":
                break
            for e in ENTS:
                out[(e, REV, m)] = base
                out[(e, EXP, m)] = -base / 2
    return out


def test_dax_parse_refresh():
    assert '>= "2025-01"' in A.dax("2025-01")
    result = {"results": [{"tables": [{"rows": [{"[l]": REV, "[e]": "02.1", "[m]": "2026-09", "[a]": 5.0},
                                                {"[l]": REV, "[e]": "02.1", "[m]": "2026-10", "[a]": None}]}]}]}
    incoming = A.parse(result)
    assert incoming[1]["amount"] == 0
    stored = [{"line": REV, "entity": "02.1", "month": m, "amount": 1.0} for m in ("2026-08", "2026-09", "2026-10")]
    assert [r["month"] for r in A.refresh(stored, incoming)] == ["2026-08", "2026-09", "2026-10"]
    assert A.refresh(stored, incoming)[1]["amount"] == 5.0


def test_baseline(actuals):
    assert F.current_month(actuals) == "2026-10"
    assert F.horizon("2026-10")[0] == "2026-10" and F.horizon("2026-10")[-1] == "2027-12"
    actuals[("02.1", REV, "2026-09")] = 1.0        # last posted month still filling: skipped by the trend
    fc = F.forecast(actuals, {})
    assert fc[("02.1", REV, "2026-11")] == pytest.approx(110)
    assert fc[("02.1", REV, "2027-03")] == pytest.approx(121)
    assert fc[("02.1", EXP, "2027-03")] == pytest.approx(-60.5)
    actuals[("03.2", REV, "2025-03")], actuals[("03.2", REV, "2026-03")] = 0.001, 50.0   # entity base ×50,000
    line = {}
    for (e, l, m), a in actuals.items():
        line[m] = line.get(m, 0) + a if l == REV else line.get(m, 0)
    g = F.yoy(line, "2026-10")
    assert abs(g - 0.1) < 0.05 and F.forecast(actuals, {})[("03.2", REV, "2027-03")] == pytest.approx((1 + g) * 50)


def test_submit(actuals):
    subs = {}
    out = F.submit([{"entity": "Pārdošana", "line": "1.1", "month": "2026-11", "amount": 600}], actuals, subs)
    assert sorted(out) == [("02.1", REV, "2026-11", 300.0), ("02.2", REV, "2026-11", 300.0)]
    out = F.submit([{"entity": "02.1", "line": REV, "month": "2026-12", "yoy": 0.5}], actuals, subs)
    assert out == [("02.1", REV, "2026-12", 150.0)]
    subs[("02.1", REV, "2026-12")] = 150.0
    assert F.forecast(actuals, subs)[("02.1", REV, "2027-12")] == pytest.approx(165)
    out = F.submit([{"entity": "TOTAL", "line": "1.1", "month": "2027-01", "amount": 363}], actuals, subs)
    assert [round(a) for *_, a in out] == [121, 121, 121]
    for bad in ({"entity": "X", "line": "1.1", "month": "2026-11", "amount": 1},
                {"entity": "02.1", "line": "9.9", "month": "2026-11", "amount": 1},
                {"entity": "02.1", "line": "1.1", "month": "2026-09", "amount": 1}):
        with pytest.raises(ValueError):
            F.submit([bad], actuals, subs)


def test_workbook(actuals):
    subs = {("02.1", REV, "2026-11"): 300.0, ("02.2", REV, "2026-11"): 300.0}
    wb = load_workbook(BytesIO(W.build(actuals, subs)))
    s = wb["Summary"]
    assert [c.value for c in s["A"][2:7]] == [REV, EXP, "1.99999 KOPĀ IEŅĒMUMI", "2.99999 KOPĀ IZDEVUMI", "3.99999 PEĻŅA"]
    assert s["B3"].value == 3600 and s["K3"].value == 330 and s["M3"].value == 710  # FY25, Sep-26 actual, Nov-26 forecast
    assert s["L3"].fill.fgColor.rgb.endswith("DCE6F1") and not s["K3"].fill.fgColor.rgb.endswith("DCE6F1")
    assert s["O3"].value == "=SUM(C3:N3)" and s["C5"].value == "=C3" and s["C7"].value == "=C5+C6"

    ws = wb[REV[:31]]
    assert ws["B1"].value == "TOTAL" and ws["F1"].value == "TOTAL Pārdošana" and ws["I1"].value == "02.1 Pārdošanas jaunas auto"
    assert ws["A3"].value == "2026-01" and ws["A26"].value == "2027-12" and ws["A27"].value == "TOTAL 2026"
    assert ws["I13"].value == 300 and ws["J13"].value == 100 and ws["K13"].value == "=IF(J13=0,0,I13/J13-1)"
    assert ws["F13"].value == "=I13+L13" and ws["B13"].value == "=F13+P13" and ws["P1"].value == "TOTAL Serviss"
    assert ws["I27"].value == "=SUM(I3:I14)" and ws["I28"].value == "=SUM(I15:I26)"
    assert ws["I13"].fill.fgColor.rgb.endswith("DCE6F1") and not ws["I11"].fill.fgColor.rgb.endswith("DCE6F1")
    assert ws.column_dimensions["I"].outlineLevel == 1 and ws.column_dimensions["E"].width == 3
